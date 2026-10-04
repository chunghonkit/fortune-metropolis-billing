"""Electricity history workbook and the bill/meter trend series."""

import asyncio
import hashlib

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from app.consumption_history import (
    history_path,
    history_series,
    load_history,
    upsert_history,
)
from app.main import download_electricity_history, electricity_history


def _bill(account, kwh, total, meters, fit_amount=0, fit_lines=None):
    return {
        'account': account,
        'kwh': kwh,
        'total_amount': total,
        'fit_amount': fit_amount,
        'fuel_amount': 12.5,
        'less_charge': 0,
        'from_date': '26-04-25',
        'to_date': '23-05-25',
        'days': 28,
        'meters': meters,
        'fit_lines': fit_lines or [],
    }


def _meter(number, kwh, previous=100, present=None):
    return {
        'meter_no': number,
        'consumption': kwh,
        'primary': kwh,
        'adjustment': 0,
        'previous': previous,
        'present': present if present is not None else previous + kwh,
        'factor': 1,
    }


def _sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_history_keeps_months_and_does_not_invent_meter_charges(tmp_path):
    april = tmp_path / '2025-04' / 'Cost Allocation.xlsx'
    april.parent.mkdir(parents=True)
    wb = Workbook()
    wb.active['A1'] = 'april-original'
    wb.save(april)
    wb.close()
    april_hash = _sha(april)

    chiller = '9024222'
    other = '9048406'
    fit_meter = '9115936'
    upsert_history(tmp_path, '2025-05', [
        _bill('55861-52267-1', 331091, 505773, [
            _meter(chiller, 130769),
            _meter(other, 82374),
        ]),
        _bill('52167-13569-2', 165060, 144373, [
            _meter(fit_meter, 6878),
        ], fit_amount=-27512, fit_lines=[{'meter': fit_meter, 'amount': 27512}]),
    ])
    upsert_history(tmp_path, '2025-06', [
        _bill('55861-52267-1', 340000, 520000, [
            _meter(chiller, 140111, previous=96504),
            _meter(other, 90000),
        ]),
    ])

    stored = load_history(tmp_path)
    may_chiller = [
        row for row in stored['meters']
        if row['billing_month'] == '2025-05' and str(row['meter_no']) == chiller
    ]
    assert len(may_chiller) == 1
    assert may_chiller[0]['kwh'] == 130769
    assert may_chiller[0]['charge'] is None
    assert may_chiller[0]['charge_source'] in (None, '')
    assert may_chiller[0]['charge'] != 505773

    fit_row = next(row for row in stored['meters'] if str(row['meter_no']) == fit_meter)
    assert fit_row['charge'] == -27512
    assert fit_row['charge_source'] == 'fit_line'

    # Reprocessing May replaces May and leaves June.
    upsert_history(tmp_path, '2025-05', [
        _bill('55861-52267-1', 111, 222, [_meter(chiller, 111)]),
    ])
    stored = load_history(tmp_path)
    may_meters = [row for row in stored['meters'] if row['billing_month'] == '2025-05']
    june_meters = [row for row in stored['meters'] if row['billing_month'] == '2025-06']
    assert [row['kwh'] for row in may_meters] == [111]
    assert sorted(row['kwh'] for row in june_meters) == [90000, 140111]
    assert _sha(april) == april_hash
    assert history_path(tmp_path) == tmp_path / 'masters' / 'electricity_history.xlsx'
    assert not (tmp_path / '2025-05' / 'out' / 'electricity_history.xlsx').exists()

    bill_trend = history_series(stored, account='55861-52267-1')
    assert bill_trend['source'] == 'electricity_history'
    assert bill_trend['months'] == ['2025-05', '2025-06']
    assert bill_trend['bill_kwh'] == [111, 340000]
    assert bill_trend['bill_charge'] == [222, 520000]

    meter_trend = history_series(stored, account='55861-52267-1', meter=chiller)
    assert meter_trend['meter_kwh'] == [111, 140111]
    assert meter_trend['meter_charge'] == [None, None]
    assert all(str(row['meter_no']) == chiller for row in meter_trend['meters'])


def test_history_download_serves_the_master(tmp_path, monkeypatch):
    monkeypatch.setenv('METROPOLIS_ROOT', str(tmp_path))
    upsert_history(tmp_path, '2025-05', [
        _bill('82805-94744-7', 49157, 70610, [_meter('9130001', 49157)]),
    ])

    payload = asyncio.run(electricity_history(account='82805-94744-7', meter='9130001'))
    assert payload['bill_kwh'] == [49157]
    assert payload['bill_charge'] == [70610]
    assert payload['meter_kwh'] == [49157]
    assert payload['meter_charge'] == [None]
    assert payload['source'] == 'electricity_history'

    response = asyncio.run(download_electricity_history())
    assert response.filename == 'electricity_history.xlsx'
    book = load_workbook(response.path)
    assert book.sheetnames[:2] == ['Bills', 'Meters']
    bills = list(book['Bills'].iter_rows(values_only=True))
    meters = list(book['Meters'].iter_rows(values_only=True))
    assert bills[1][0] == '2025-05'
    assert bills[1][1] == '82805-94744-7'
    assert bills[1][5] == 49157
    assert bills[1][6] == 70610
    assert meters[1][2] == '9130001'
    assert meters[1][6] == 49157
    assert meters[1][9] is None
    book.close()

    empty = tmp_path / 'empty'
    monkeypatch.setenv('METROPOLIS_ROOT', str(empty))
    with pytest.raises(HTTPException) as absent:
        asyncio.run(download_electricity_history())
    assert absent.value.status_code == 404
