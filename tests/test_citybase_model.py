"""
Citybase master formulas, checked against the April workbook's own cache.

These tests do not freeze May Cost Sheet dollars. They check that the
check-meter rule and the AC DEPT leaf rows reproduce the master.
"""

from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.citybase_model import CitybaseModel, canonical_centre_key

MASTER = Path('/home/ubuntu/.cursor/projects/workspace/uploads/april-cost-allocation_f85f.xlsx')


pytestmark = pytest.mark.skipif(not MASTER.exists(), reason='April master not in this workspace')


def test_shared_label_keeps_every_segment():
    assert canonical_centre_key('Hotel / Commercial / SA(11)') == 'Hotel / Commercial / SA'
    assert canonical_centre_key('Hotel/Commercial (10)') == 'Hotel / Commercial'
    assert canonical_centre_key('SA/Commercial (16)') == 'SA / Commercial'
    assert canonical_centre_key('Commercial Common (2)') == 'Commercial Common'
    assert canonical_centre_key('AC - Commercial Air Conditioning') == 'AC'
    assert canonical_centre_key('OC -Office Common') == 'OC'
    assert canonical_centre_key('O') == 'O'
    assert canonical_centre_key('C') == 'C'


def _april_bills_and_log():
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    elect = wb['Elect Charge']
    # Column BC is the April 2025 column named by Allocation formulas.
    def bc(row):
        return elect.cell(row=row, column=55).value

    chiller_kwh = {
        '9048406': bc(28),
        '9026169': bc(29),
        '9024222': bc(30),
        '9026183': bc(31),
    }
    fit_kwh = {
        '9133719': bc(36),
        '9134412': bc(37),
        '9144186': bc(38),
    }
    # Direct elect-row amounts, keyed by the account Max Demand names.
    direct = {
        '35204-69738-4': bc(11),
        '40722-61440-7': bc(12),
        '00776-78552-1': bc(14),
        '13639-58422-3': bc(15),
        '79292-23337-6': bc(16),
        '24096-78457-6': bc(17),
        '08731-83914-5': bc(18),
        '97968-02236-6': bc(19),
        '23529-59279-9': bc(20),
        '88931-57029-6': bc(22),
        '70873-85471-3': bc(23),
        '72399-00664-9': bc(24),
    }
    wb.close()

    bills = [{
        'account': '55861-52267-1',
        'total_amount': bc_amount_from(chiller_kwh, direct, kind='chiller'),
        'meters': [{'meter_no': meter, 'consumption': kwh} for meter, kwh in chiller_kwh.items()],
    }]
    # Re-read bill total from the sheet (row 34) rather than reconstructing.
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    elect = wb['Elect Charge']
    bills[0]['total_amount'] = elect.cell(row=34, column=55).value
    fit_total = elect.cell(row=41, column=55).value
    for account, amount in direct.items():
        bills.append({'account': account, 'total_amount': amount, 'meters': []})
    bills.append({
        'account': '52167-13569-2',
        'total_amount': fit_total,
        'meters': [{'meter_no': meter, 'consumption': kwh} for meter, kwh in fit_kwh.items()],
    })
    # April dial on the check meter: 96323.1 − 95966.6 = 356.5
    meter_log = {'meter_no': '6681757', 'previous': 95966.6, 'present': 96323.1}
    ac = wb['AC DEPT']
    cached = {}
    for row in range(4, 70):
        label = ac.cell(row=row, column=15).value
        amount = ac.cell(row=row, column=3).value
        if label is not None and isinstance(amount, (int, float)):
            cached[row] = (canonical_centre_key(str(label)), amount)
    c73 = ac['C73'].value
    c74 = ac['C74'].value
    wb.close()
    return bills, meter_log, cached, c73, c74


def bc_amount_from(*_args, **_kwargs):
    return 0


def test_april_check_meter_split_matches_master_columns():
    bills, meter_log, _cached, c73, c74 = _april_bills_and_log()
    model = CitybaseModel(str(MASTER))
    result = model.allocate(bills, meter_log)

    ac = sum(a['amount'] for a in result['allocations']
             if a['account'] == '55861-52267-1' and a['centre'] == 'AC')
    sw = sum(a['amount'] for a in result['allocations']
             if a['account'] == '55861-52267-1' and a['centre'] == 'SW')
    assert abs(ac - c73) < 0.02
    assert abs(sw - c74) < 0.02
    assert abs((ac / (ac + sw)) - (c73 / (c73 + c74))) < 1e-6
    # April spot-check quoted by the workbook's own C73/C74.
    assert abs((ac / (ac + sw)) - 0.778194) < 0.0001


def test_april_leaf_rows_match_cached_master():
    bills, meter_log, cached, _c73, _c74 = _april_bills_and_log()
    model = CitybaseModel(str(MASTER))
    result = model.allocate(bills, meter_log)

    # Sum of leaves by centre, from the master's cached column C, using the
    # same leaf-row list the model read from the Total formula.
    expected = {}
    for row in model.leaf_rows:
        if row not in cached:
            continue
        label, amount = cached[row]
        expected[label] = expected.get(label, 0) + amount

    actual = {}
    for alloc in result['allocations']:
        if alloc['centre'] == 'FC':
            continue
        actual[alloc['centre']] = actual.get(alloc['centre'], 0) + alloc['amount']

    assert 'Hotel / Commercial / SA' in actual
    assert 'Commercial Common' in actual
    assert 'Hotel / L8 Premises' in actual
    for label, amount in expected.items():
        assert abs(actual.get(label, 0) - amount) < 0.05, (
            f'{label}: model {actual.get(label, 0)} vs master {amount}'
        )


def test_may_check_meter_moves_o7_without_using_building_total():
    """
    O7 is (delta × 160) / kWh of meter 9024222 only.

    Dividing the dial delta by the building total is the bug that produced
    AC 99.95% / SW 0.05%. This uses a made-up May dial and kWh, not the
    golden Cost Sheet.
    """
    bills, _log, _cached, _c73, _c74 = _april_bills_and_log()
    # Replace only the 9024222 kWh and the dial. Other April inputs stay,
    # so the assertion is about the formula, not the May golden file.
    for meter in bills[0]['meters']:
        if meter['meter_no'] == '9024222':
            meter['consumption'] = 130752
    meter_log = {'meter_no': '6681757', 'previous': 96323.1, 'present': 96504.6}
    # delta 181.5 × 160 / 130752 ≈ 0.2221
    model = CitybaseModel(str(MASTER))
    # Reach the percentage table directly.
    o7 = None
    percentages = model._percentages(Decimal('181.5'), Decimal(130752))
    for row, line in model.alloc_lines.items():
        if line['kind'] == 'check_sw':
            o7 = percentages[row]
    assert o7 is not None
    assert abs(float(o7) - (181.5 * 160 / 130752)) < 1e-12
    assert abs(float(o7) - 0.2221) < 0.0001

    result = model.allocate(bills, meter_log)
    ac = sum(a['amount'] for a in result['allocations']
             if a['account'] == '55861-52267-1' and a['centre'] == 'AC')
    sw = sum(a['amount'] for a in result['allocations']
             if a['account'] == '55861-52267-1' and a['centre'] == 'SW')
    # A smaller SW share of meter 9024222 moves the account off April's 77.82 / 22.18.
    assert ac / (ac + sw) > 0.7782
    assert sw / (ac + sw) < 0.2218


def test_chinese_bill_meters_feed_check_meter_denominator():
    """9024222 units come from the 電錶號碼 block, not the building total."""
    from app.clp_parser import _extract_chinese_meter_units

    text = """
    電錶號碼
    度數
    總數
    9048406
    80,000.00
    9026169
    50,000.00
    9024222
    130,752.00
    9026183
    1,000.00
    總用電度數
    261752.00
    """
    parsed = {'meters': []}
    _extract_chinese_meter_units(text, parsed)
    by_meter = {m['meter_no']: m['consumption'] for m in parsed['meters']}
    assert by_meter['9024222'] == 130752
    assert set(by_meter) == {'9048406', '9026169', '9024222', '9026183'}

    bills = [{
        'account': '55861-52267-1',
        'total_amount': 500000,
        'meters': parsed['meters'],
    }]
    meter_log = {'meter_no': '6681757', 'previous': 96323.1, 'present': 96504.6}
    result = CitybaseModel(str(MASTER)).allocate(bills, meter_log)
    ac = sum(a['amount'] for a in result['allocations']
             if a['account'] == '55861-52267-1' and a['centre'] == 'AC')
    sw = sum(a['amount'] for a in result['allocations']
             if a['account'] == '55861-52267-1' and a['centre'] == 'SW')
    # Frozen April Column C is 77.82 / 22.18. Live O7 ≈ 0.2221 moves this.
    assert abs((ac / (ac + sw)) - 0.778194) > 0.01
    assert abs(ac + sw - 500000) < 0.05


def test_office_chillers_92_8_and_carpark_shared_rows_balance():
    model = CitybaseModel(str(MASTER))
    bills = [
        {'account': '13639-58422-3', 'total_amount': 10000, 'meters': []},
        {'account': '79292-23337-6', 'total_amount': 20000, 'meters': []},
        {'account': '35204-69738-4', 'total_amount': 50000, 'meters': []},
    ]
    result = model.allocate(bills, {'meter_no': '6681757', 'previous': 1, 'present': 1})

    def total(account, centre=None):
        return sum(
            a['amount'] for a in result['allocations']
            if a['account'] == account and (centre is None or a['centre'] == centre)
        )

    assert abs(total('13639-58422-3', 'AO') - 9200) < 0.02
    assert abs(total('13639-58422-3', 'OC') - 800) < 0.02
    assert abs(total('79292-23337-6', 'AO') - 18400) < 0.02
    assert abs(total('79292-23337-6', 'OC') - 1600) < 0.02
    carpark_rows = [a for a in result['allocations'] if a['account'] == '35204-69738-4']
    assert abs(sum(a['amount'] for a in carpark_rows) - 50000) < 0.05
    centres = {a['centre'] for a in carpark_rows}
    assert 'CP' in centres
    assert 'Hotel / L8 Premises' in centres
    assert 'Commercial Common' in centres
