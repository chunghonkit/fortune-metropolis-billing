"""Dashboard downloads and chart series for a processed month."""

import asyncio
import json

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from app.dashboard_data import build_month_report, chart_series, save_month_report
from app.main import (
    dashboard_charts,
    dashboard_month,
    download_output_workbook,
    processed_months,
)


def _workbook(path, marker):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.active['A1'] = marker
    wb.save(path)
    wb.close()


def _bill(account, kwh, total, fit=0):
    return {
        'account': account,
        'kwh': kwh,
        'total_amount': total,
        'fit_amount': fit,
        'meters': [{'meter_no': '9999999', 'consumption': 1, 'charge': 99999}],
    }


def _allocation(account, centre, percentage, amount, base):
    return {
        'account': account,
        'centre': centre,
        'percentage': percentage,
        'amount': amount,
        'allocation_base': base,
        'bill_total': base,
        'kwh': 0,
    }


def test_chart_series_follow_bills_and_allocation():
    may_kwh = 1234
    june_kwh = 5678
    may_amount = 87.65
    june_amount = 43.21
    percentage = 0.375
    may = build_month_report(
        '2025-05',
        [_bill('52167-13569-2', may_kwh, 100, fit=-20)],
        {'allocations': [
            _allocation('52167-13569-2', 'AC', percentage, may_amount, 120),
            _allocation('52167-13569-2', 'SW', 1 - percentage, 32.35, 120),
        ]},
    )
    june = build_month_report(
        '2025-06',
        [_bill('82805-94744-7', june_kwh, 50)],
        {'allocations': [
            _allocation('82805-94744-7', 'FC', 1, june_amount, 50),
        ]},
    )

    assert may['accounts'][0]['allocation_base'] == 120
    assert may['accounts'][0]['kwh'] == may_kwh

    both = chart_series([may, june])
    assert both['source'] == 'parsed_bills+allocation'
    assert both['months'] == ['2025-05', '2025-06']
    assert both['kwh'] == [may_kwh, june_kwh]
    assert both['cost'] == pytest.approx([may_amount + 32.35, june_amount])
    # A meter charge printed nowhere on the bill is not the centre cost.
    assert all(row['amount'] != 99999 for row in both['rows'])

    ac = chart_series([may, june], centre='AC')
    assert ac['months'] == ['2025-05', '2025-06']
    assert ac['kwh'] == pytest.approx([may_kwh * percentage, 0])
    assert ac['cost'] == pytest.approx([may_amount, 0])
    assert ac['by_centre'] == pytest.approx([
        {'name': 'AC', 'kwh': may_kwh * percentage, 'cost': may_amount}
    ])

    one_month = chart_series([may, june], months=['2025-06'], account='82805-94744-7')
    assert one_month['months'] == ['2025-06']
    assert one_month['kwh'] == [june_kwh]
    assert one_month['cost'] == [june_amount]
    assert [row['centre'] for row in one_month['rows']] == ['FC']


def _run(coro):
    return asyncio.run(coro)


def test_processed_month_workbooks_download(tmp_path, monkeypatch):
    monkeypatch.setenv('METROPOLIS_ROOT', str(tmp_path))
    root = tmp_path
    month = '2025-05'
    allocation = root / month / 'out' / 'Cost Allocation - Electricity 2007-2.xlsx'
    sheet = root / month / 'out' / 'Cost Sheet-2025-05.xlsx'
    _workbook(allocation, 'allocation-marker')
    _workbook(sheet, 'sheet-marker')
    # A second Cost Allocation must not hide the file the report names.
    _workbook(root / month / 'out' / 'Cost Allocation other.xlsx', 'other-marker')

    report = save_month_report(
        root,
        month,
        [_bill('97968-02236-6', 2320, 27503)],
        {'allocations': [_allocation('97968-02236-6', 'DC', 1, 27503, 27503)], 'summary': {}},
        {'cost_allocation': allocation, 'cost_sheet': sheet},
    )
    saved = json.loads(report.read_text(encoding='utf-8'))
    assert saved['files']['cost_allocation'] == allocation.name
    assert saved['accounts'][0]['kwh'] == 2320

    listed = _run(processed_months())
    assert listed['months'] == [{
        'month': month,
        'has_report': True,
        'cost_allocation': True,
        'cost_sheet': True,
    }]

    dashboard = _run(dashboard_month(month))
    assert dashboard['accounts'][0]['kwh'] == 2320
    assert dashboard['allocations'][0]['amount'] == 27503

    body = _run(dashboard_charts(months=month, centre='DC'))
    assert body['kwh'] == [2320]
    assert body['cost'] == [27503]
    assert body['source'] == 'parsed_bills+allocation'

    for kind, marker, filename in (
        ('cost-allocation', 'allocation-marker', allocation.name),
        ('cost-sheet', 'sheet-marker', sheet.name),
    ):
        response = _run(download_output_workbook(month, kind))
        assert 'spreadsheet' in response.media_type
        assert response.filename == filename
        assert load_workbook(response.path).active['A1'].value == marker

    with pytest.raises(HTTPException) as missing:
        _run(download_output_workbook('2025-06', 'cost-allocation'))
    assert missing.value.status_code == 404
    with pytest.raises(HTTPException) as bad_kind:
        _run(download_output_workbook('2025-05', 'not-a-workbook'))
    assert bad_kind.value.status_code == 400
    with pytest.raises(HTTPException) as bad_month:
        _run(download_output_workbook('../etc', 'cost-allocation'))
    assert bad_month.value.status_code == 400
    with pytest.raises(HTTPException) as bad_charts:
        _run(dashboard_charts(months='August'))
    assert bad_charts.value.status_code == 400
