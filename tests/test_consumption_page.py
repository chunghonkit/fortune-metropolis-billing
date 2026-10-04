"""Consumption page buttons and one-item charts from the history workbook."""

import asyncio
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.consumption_history import history_path
from app.dashboard_data import save_month_report
from app.main import history_item_buttons, history_item_chart


def _run(coro):
    return asyncio.run(coro)


def _bill(account, kwh, total, meters, fit_amount=0, fit_lines=None):
    return {
        'account': account,
        'kwh': kwh,
        'total_amount': total,
        'fit_amount': fit_amount,
        'meters': meters,
        'fit_lines': fit_lines or [],
    }


def _meter(number, kwh):
    return {'meter_no': number, 'consumption': kwh, 'primary': kwh, 'adjustment': 0}


def _allocation(account, centre, percentage, amount, base):
    return {
        'account': account,
        'centre': centre,
        'percentage': percentage,
        'amount': amount,
        'allocation_base': base,
    }


def _seed(tmp_path):
    chiller = '9024222'
    fit_meter = '9115936'
    retail = '55861-52267-1'
    fit_account = '52167-13569-2'
    percentage = 0.6
    save_month_report(
        tmp_path,
        '2025-05',
        [
            _bill(retail, 1000, 400, [_meter(chiller, 250)]),
            _bill(fit_account, 80, 30, [_meter(fit_meter, 80)], fit_amount=-10, fit_lines=[{'meter': fit_meter, 'amount': 10}]),
        ],
        {'allocations': [
            _allocation(retail, 'AC', percentage, 240, 400),
            _allocation(retail, 'SW', 1 - percentage, 160, 400),
            _allocation(fit_account, 'C', 1, 40, 40),
        ]},
    )
    save_month_report(
        tmp_path,
        '2025-06',
        [_bill(retail, 1500, 500, [_meter(chiller, 300)])],
        {'allocations': [
            _allocation(retail, 'AC', percentage, 300, 500),
            _allocation(retail, 'SW', 1 - percentage, 200, 500),
        ]},
    )
    return retail, chiller, fit_account, fit_meter, percentage


def test_buttons_switch_and_chart_reads_history_file(tmp_path, monkeypatch):
    monkeypatch.setenv('METROPOLIS_ROOT', str(tmp_path))
    retail, chiller, fit_account, fit_meter, percentage = _seed(tmp_path)

    bills = _run(history_item_buttons('bill'))
    meters = _run(history_item_buttons('meter'))
    centres = _run(history_item_buttons('centre'))
    assert bills['kind'] == 'bill'
    assert meters['kind'] == 'meter'
    assert centres['kind'] == 'centre'
    assert [item['id'] for item in bills['items']] == sorted([retail, fit_account])
    assert [item['id'] for item in meters['items']] == [
        f'{retail}|{chiller}',
        f'{fit_account}|{fit_meter}',
    ]
    assert [item['label'] for item in meters['items']] == sorted([chiller, fit_meter])
    assert [item['id'] for item in centres['items']] == ['AC', 'C', 'SW']
    for payload in (bills, meters, centres):
        assert payload['history'] is True
        assert all(set(item) == {'id', 'label'} for item in payload['items'])

    page = Path('static/index.html').read_text(encoding='utf-8')
    assert 'id="itemButtons"' in page
    assert "setItemKind('meter')" in page
    assert "setItemKind('centre')" in page
    assert 'id="consumptionTable"' not in page
    assert 'id="allocationTable"' not in page
    assert 'id="historyMeterTable"' not in page
    assert '<table' not in page
    assert 'id="dlAllocation"' in page
    assert 'id="dlSheet"' in page
    assert 'id="dlHistory"' in page

    book = load_workbook(history_path(tmp_path))
    bill_rows = list(book['Bills'].iter_rows(values_only=True))
    header = bill_rows[0]
    kwh_col = header.index('kwh')
    charge_col = header.index('total_amount')
    stored = {
        (row[0], row[1]): (row[kwh_col], row[charge_col])
        for row in bill_rows[1:]
    }
    series = _run(history_item_chart('bill', retail))
    assert series['source'] == 'electricity_history'
    assert series['months'] == ['2025-05', '2025-06']
    assert series['kwh'] == [stored[('2025-05', retail)][0], stored[('2025-06', retail)][0]]
    assert series['cost'] == [stored[('2025-05', retail)][1], stored[('2025-06', retail)][1]]
    assert 'bills' not in series
    assert 'meters' not in series

    # A later edit of the workbook is what the chart reads.
    for row in book['Bills'].iter_rows(min_row=2):
        if row[0].value == '2025-06' and row[1].value == retail:
            row[kwh_col].value = 2222
    book.save(history_path(tmp_path))
    book.close()
    reread = _run(history_item_chart('bill', retail))
    assert reread['kwh'] == [1000, 2222]
    assert reread['cost'] == [400, 500]

    meter = _run(history_item_chart('meter', f'{retail}|{chiller}'))
    assert meter['source'] == 'electricity_history'
    assert meter['kwh'] == [250, 300]
    assert meter['cost'] == [None, None]
    assert meter['cost'] != [400, 500]
    assert meter['note']

    fit = _run(history_item_chart('meter', f'{fit_account}|{fit_meter}'))
    assert fit['kwh'] == [80]
    assert fit['cost'] == [-10]

    centre = _run(history_item_chart('centre', 'AC'))
    assert centre['months'] == ['2025-05', '2025-06']
    assert centre['kwh'] == pytest.approx([1000 * percentage, 2222 * percentage])
    assert centre['cost'] == [240, 300]
    assert centre['source'] == 'electricity_history'

    with pytest.raises(HTTPException) as bad_kind:
        _run(history_item_buttons('voucher'))
    assert bad_kind.value.status_code == 400
    with pytest.raises(HTTPException) as missing:
        _run(history_item_chart('bill', '00000-00000-0'))
    assert missing.value.status_code == 404
