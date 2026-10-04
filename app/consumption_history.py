"""
Long-lived electricity history.

One workbook under METROPOLIS_ROOT/masters stores every processed month.
Reprocessing a month replaces that month's rows and leaves the others.

Bill rows are the parsed account kWh and the charges printed on the bill.
Meter rows are each meter's kWh. A meter charge is stored only when the
bill prints one (a FiT line). The account total is not spread across meters.
"""

from pathlib import Path
from typing import Dict, Iterable, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HISTORY_NAME = 'electricity_history.xlsx'

BILL_COLUMNS = (
    'billing_month',
    'account',
    'from_date',
    'to_date',
    'days',
    'kwh',
    'total_amount',
    'fit_amount',
    'fuel_amount',
    'less_charge',
)

METER_COLUMNS = (
    'billing_month',
    'account',
    'meter_no',
    'previous',
    'present',
    'factor',
    'kwh',
    'primary',
    'adjustment',
    'charge',
    'charge_source',
)


def history_path(metropolis_root: Path) -> Path:
    return Path(metropolis_root) / 'masters' / HISTORY_NAME


def _printed_meter_charges(bill: Dict) -> Dict[str, float]:
    """FiT amounts the bill prints against a meter id. Nothing else."""
    charges: Dict[str, float] = {}
    for line in bill.get('fit_lines') or []:
        meter = str(line.get('meter') or '').strip()
        if not meter or line.get('amount') is None:
            continue
        charges[meter] = charges.get(meter, 0.0) - abs(float(line['amount']))
    return charges


def bill_rows_for_month(billing_month: str, parsed_bills: Iterable[Dict]) -> List[Dict]:
    rows = []
    for bill in parsed_bills or []:
        account = bill.get('account')
        if not account:
            continue
        rows.append({
            'billing_month': billing_month,
            'account': account,
            'from_date': bill.get('from_date') or '',
            'to_date': bill.get('to_date') or '',
            'days': bill.get('days'),
            'kwh': float(bill.get('kwh') or 0),
            'total_amount': float(bill.get('total_amount') or 0),
            'fit_amount': float(bill.get('fit_amount') or 0),
            'fuel_amount': float(bill.get('fuel_amount') or 0),
            'less_charge': float(bill.get('less_charge') or 0),
        })
    rows.sort(key=lambda row: row['account'])
    return rows


def meter_rows_for_month(billing_month: str, parsed_bills: Iterable[Dict]) -> List[Dict]:
    rows = []
    for bill in parsed_bills or []:
        account = bill.get('account')
        if not account:
            continue
        printed = _printed_meter_charges(bill)
        for meter in bill.get('meters') or []:
            number = str(meter.get('meter_no') or '').strip()
            if not number:
                continue
            charge = printed.get(number)
            rows.append({
                'billing_month': billing_month,
                'account': account,
                'meter_no': number,
                'previous': meter.get('previous'),
                'present': meter.get('present'),
                'factor': meter.get('factor'),
                'kwh': float(meter.get('consumption') if meter.get('consumption') is not None else meter.get('kwh') or 0),
                'primary': meter.get('primary'),
                'adjustment': meter.get('adjustment'),
                'charge': charge,
                'charge_source': 'fit_line' if charge is not None else '',
            })
    rows.sort(key=lambda row: (row['account'], row['meter_no']))
    return rows


def _read_sheet(ws) -> List[Dict]:
    raw = list(ws.iter_rows(values_only=True))
    if not raw:
        return []
    header = [str(cell) if cell is not None else '' for cell in raw[0]]
    rows = []
    for values in raw[1:]:
        if not values or not any(value is not None and value != '' for value in values):
            continue
        rows.append({key: values[index] if index < len(values) else None for index, key in enumerate(header)})
    return rows


def load_history(metropolis_root: Path) -> Dict[str, List[Dict]]:
    path = history_path(metropolis_root)
    if not path.is_file():
        return {'bills': [], 'meters': []}
    wb = load_workbook(path, data_only=False)
    try:
        bills = _read_sheet(wb['Bills']) if 'Bills' in wb.sheetnames else []
        meters = _read_sheet(wb['Meters']) if 'Meters' in wb.sheetnames else []
    finally:
        wb.close()
    return {'bills': bills, 'meters': meters}


def _write_sheet(ws, columns, rows):
    ws.append(list(columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(column) for column in columns])
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = max(12, len(str(column[0].value or '')) + 2)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def _save(path: Path, bills: List[Dict], meters: List[Dict]) -> None:
    wb = Workbook()
    bills_sheet = wb.active
    bills_sheet.title = 'Bills'
    _write_sheet(bills_sheet, BILL_COLUMNS, bills)
    meters_sheet = wb.create_sheet('Meters')
    _write_sheet(meters_sheet, METER_COLUMNS, meters)
    note = wb.create_sheet('Read me')
    note['A1'] = (
        'Bills are one row per account per month: parsed kWh and the charges printed on that bill. '
        'Meters are one row per meter: kWh from the register. '
        'charge is filled only when the bill prints an amount for that meter (a FiT line). '
        'The account total is not divided across meters.'
    )
    note.column_dimensions['A'].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def upsert_history(metropolis_root: Path, billing_month: str, parsed_bills: Iterable[Dict]) -> Path:
    """Replace this month's rows. Earlier and later months stay."""
    from app.dashboard_data import valid_month
    if not valid_month(billing_month):
        raise ValueError(f'billing month must be YYYY-MM, got {billing_month!r}')
    path = history_path(metropolis_root)
    existing = load_history(metropolis_root)
    bills = [row for row in existing['bills'] if row.get('billing_month') != billing_month]
    meters = [row for row in existing['meters'] if row.get('billing_month') != billing_month]
    bills.extend(bill_rows_for_month(billing_month, parsed_bills))
    meters.extend(meter_rows_for_month(billing_month, parsed_bills))
    bills.sort(key=lambda row: (str(row.get('billing_month') or ''), str(row.get('account') or '')))
    meters.sort(key=lambda row: (
        str(row.get('billing_month') or ''),
        str(row.get('account') or ''),
        str(row.get('meter_no') or ''),
    ))
    _save(path, bills, meters)
    return path


def history_series(
    records: Dict[str, List[Dict]],
    account: Optional[str] = None,
    meter: Optional[str] = None,
) -> Dict:
    """
    Trend series read back from the history workbook.

    Bill kWh and bill charge are the stored account rows. Meter kWh is the
    stored register kWh. meter_charge is null in a month where that meter
    has no printed charge.
    """
    bills = list(records.get('bills') or [])
    meters = list(records.get('meters') or [])
    if account:
        bills = [row for row in bills if row.get('account') == account]
        meters = [row for row in meters if row.get('account') == account]
    if meter:
        meters = [row for row in meters if str(row.get('meter_no')) == str(meter)]

    months = sorted({
        str(row.get('billing_month'))
        for row in bills + meters
        if row.get('billing_month')
    })
    bill_kwh = []
    bill_charge = []
    meter_kwh = []
    meter_charge = []
    for month in months:
        month_bills = [row for row in bills if str(row.get('billing_month')) == month]
        month_meters = [row for row in meters if str(row.get('billing_month')) == month]
        bill_kwh.append(sum(float(row.get('kwh') or 0) for row in month_bills))
        bill_charge.append(sum(float(row.get('total_amount') or 0) for row in month_bills))
        meter_kwh.append(sum(float(row.get('kwh') or 0) for row in month_meters))
        printed = [float(row['charge']) for row in month_meters if row.get('charge') is not None]
        meter_charge.append(sum(printed) if printed else None)
    return {
        'months': months,
        'bill_kwh': bill_kwh,
        'bill_charge': bill_charge,
        'meter_kwh': meter_kwh,
        'meter_charge': meter_charge,
        'bills': bills,
        'meters': meters,
        'source': 'electricity_history',
    }


def _meter_id(account: str, meter_no: str) -> str:
    return f'{account}|{meter_no}'


def history_items(records: Dict[str, List[Dict]], reports: Iterable[Dict], kind: str) -> List[Dict]:
    """
    One button per bill, meter, or cost centre.

    The list is ids and labels only. kWh and charges stay in the workbook.
    """
    if kind == 'bill':
        accounts = sorted({
            row.get('account')
            for row in records.get('bills') or []
            if row.get('account')
        })
        return [{'id': account, 'label': account} for account in accounts]
    if kind == 'meter':
        seen = set()
        meters = []
        for row in records.get('meters') or []:
            meter_no = str(row.get('meter_no') or '').strip()
            account = row.get('account') or ''
            if not meter_no:
                continue
            item_id = _meter_id(account, meter_no)
            if item_id in seen:
                continue
            seen.add(item_id)
            meters.append({'id': item_id, 'label': meter_no, 'account': account})
        counts: Dict[str, int] = {}
        for item in meters:
            counts[item['label']] = counts.get(item['label'], 0) + 1
        buttons = []
        for item in sorted(meters, key=lambda item: (item['label'], item['account'])):
            label = item['label']
            if counts[label] > 1 and item['account']:
                label = f"{label} · {item['account']}"
            buttons.append({'id': item['id'], 'label': label})
        return buttons
    if kind == 'centre':
        centres = sorted({
            (line.get('centre') or '')
            for report in reports or []
            for line in report.get('allocations') or []
            if line.get('centre')
        })
        return [{'id': centre, 'label': centre} for centre in centres]
    raise ValueError('kind must be bill, meter, or centre')


def _rows_for_month(rows: List[Dict], month: str) -> List[Dict]:
    return [row for row in rows if str(row.get('billing_month')) == month]


def _history_bill_kwh(records: Dict[str, List[Dict]], month: str, account: str) -> float:
    total = 0.0
    for row in records.get('bills') or []:
        if str(row.get('billing_month')) == month and row.get('account') == account:
            total += float(row.get('kwh') or 0)
    return total


def item_chart_series(
    records: Dict[str, List[Dict]],
    reports: Iterable[Dict],
    kind: str,
    item_id: str,
) -> Optional[Dict]:
    """
    kWh and cost for one clicked item, across the months saved for it.

    Bill and meter numbers are read from the history workbook. A meter cost
    is null when that month has no printed charge. A cost centre's kWh is
    each account's history kWh times that month's saved live percent. The
    centre cost is the saved allocation amount, not a meter charge.
    """
    if kind == 'bill':
        rows = [row for row in records.get('bills') or [] if row.get('account') == item_id]
        if not rows:
            return None
        months = sorted({str(row.get('billing_month')) for row in rows if row.get('billing_month')})
        kwh = []
        cost = []
        for month in months:
            month_rows = _rows_for_month(rows, month)
            kwh.append(sum(float(row.get('kwh') or 0) for row in month_rows))
            cost.append(sum(float(row.get('total_amount') or 0) for row in month_rows))
        return _chart(kind, item_id, item_id, months, kwh, cost, None)

    if kind == 'meter':
        meters = list(records.get('meters') or [])
        if '|' in item_id:
            account, meter_no = item_id.split('|', 1)
            rows = [
                row for row in meters
                if row.get('account') == account and str(row.get('meter_no')) == meter_no
            ]
            label = meter_no
        else:
            rows = [row for row in meters if str(row.get('meter_no')) == item_id]
            label = item_id
        if not rows:
            return None
        months = sorted({str(row.get('billing_month')) for row in rows if row.get('billing_month')})
        kwh = []
        cost = []
        for month in months:
            month_rows = _rows_for_month(rows, month)
            kwh.append(sum(float(row.get('kwh') or 0) for row in month_rows))
            printed = [float(row['charge']) for row in month_rows if row.get('charge') is not None]
            cost.append(sum(printed) if printed else None)
        note = None
        if all(value is None for value in cost):
            note = 'No charge is printed for this meter. The account total stays on the bill.'
        return _chart(kind, item_id, label, months, kwh, cost, note)

    if kind == 'centre':
        selected = []
        for report in reports or []:
            lines = [
                line for line in report.get('allocations') or []
                if line.get('centre') == item_id
            ]
            if lines:
                selected.append((report.get('billing_month'), lines))
        selected = [pair for pair in selected if pair[0]]
        if not selected:
            return None
        selected.sort(key=lambda pair: pair[0])
        months = []
        kwh = []
        cost = []
        for month, lines in selected:
            months.append(month)
            attributed = 0.0
            amount = 0.0
            for line in lines:
                attributed += _history_bill_kwh(records, month, line.get('account') or '') * float(line.get('percentage') or 0)
                amount += float(line.get('amount') or 0)
            kwh.append(attributed)
            cost.append(amount)
        return _chart(kind, item_id, item_id, months, kwh, cost, None)

    raise ValueError('kind must be bill, meter, or centre')


def _chart(kind, item_id, label, months, kwh, cost, note) -> Dict:
    return {
        'kind': kind,
        'id': item_id,
        'label': label,
        'months': months,
        'kwh': kwh,
        'cost': cost,
        'source': 'electricity_history',
        'note': note,
    }


def history_choices(records: Dict[str, List[Dict]]) -> Dict[str, List[Dict]]:
    accounts = []
    seen_accounts = set()
    for row in records.get('bills') or []:
        account = row.get('account')
        if account and account not in seen_accounts:
            seen_accounts.add(account)
            accounts.append(account)
    meters = []
    seen_meters = set()
    for row in records.get('meters') or []:
        key = (row.get('account'), str(row.get('meter_no')))
        if not key[1] or key in seen_meters:
            continue
        seen_meters.add(key)
        meters.append({'account': key[0], 'meter_no': key[1]})
    return {
        'accounts': sorted(accounts),
        'meters': sorted(meters, key=lambda item: (item['account'] or '', item['meter_no'])),
    }
