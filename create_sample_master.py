"""
Create sample master workbook from allocation_config.json for testing/defaults.
"""

import json
import openpyxl
from openpyxl.styles import Font, PatternFill

def create_sample_master():
    # Load existing allocation config
    with open('allocation_config.json', 'r') as f:
        config = json.load(f)
    
    wb = openpyxl.Workbook()
    
    # Create Allocation sheet
    if 'Sheet' in wb.sheetnames:
        ws = wb['Sheet']
        ws.title = 'Allocation'
    else:
        ws = wb.create_sheet('Allocation')
    
    # Headers
    ws['A1'] = 'Account/Meter'
    ws['A1'].font = Font(bold=True)
    ws['A1'].fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    
    # Collect all unique centres
    centres = set()
    for account_data in config.values():
        for alloc in account_data.get('allocations', []):
            centres.add(alloc['centre'])
    
    centres = sorted(centres)
    
    # Write centre headers
    for idx, centre in enumerate(centres, start=2):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws[f'{col_letter}1'] = centre
        ws[f'{col_letter}1'].font = Font(bold=True)
        ws[f'{col_letter}1'].fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    
    # Write data rows
    row = 2
    for account_key, account_data in sorted(config.items()):
        account = account_data.get('account', account_key)
        ws[f'A{row}'] = account
        
        # Create percentage map
        pct_map = {}
        for alloc in account_data.get('allocations', []):
            pct_map[alloc['centre']] = alloc['perc']
        
        # Write percentages
        for idx, centre in enumerate(centres, start=2):
            col_letter = openpyxl.utils.get_column_letter(idx)
            if centre in pct_map:
                # Write as percentage (0.xx format)
                ws[f'{col_letter}{row}'] = pct_map[centre]
                ws[f'{col_letter}{row}'].number_format = '0.0000'
        
        row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 20
    for idx in range(2, len(centres) + 2):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = 12
    
    # Create AC DEPT sheet (copy of Allocation for now)
    ws_ac = wb.create_sheet('AC DEPT')
    for row in ws.iter_rows():
        for cell in row:
            new_cell = ws_ac[cell.coordinate]
            new_cell.value = cell.value
            if cell.font:
                new_cell.font = cell.font.copy()
            if cell.fill:
                new_cell.fill = cell.fill.copy()
    
    # Save
    wb.save('data/masters/cost_allocation_master.xlsx')
    print("Created sample master at data/masters/cost_allocation_master.xlsx")

if __name__ == '__main__':
    create_sample_master()
